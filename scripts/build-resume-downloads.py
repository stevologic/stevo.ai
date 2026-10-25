"""Create the resume downloads from verified master text.

Writes both public/resume-files/Stephen-Abbott-Resume.docx and
Stephen-Abbott-Resume.pdf from content/resume-download.json. Requires
python-docx, pypdf, and reportlab (see scripts/requirements-resume-downloads.txt).

After this script, run scripts/verify-resume-downloads.py --finalize.
"""

from pathlib import Path
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

ROOT = Path(__file__).resolve().parents[1]
source = json.loads((ROOT / 'content/resume-download.json').read_text(encoding='utf-8'))
document = Document()
for style in document.styles:
    properties = style._element.find(qn('w:pPr'))
    if properties is not None:
        for border in list(properties.findall(qn('w:pBdr'))):
            properties.remove(border)
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.65)
section.footer_distance = Inches(0.3)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'List Bullet'):
    style = document.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Calibri')
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Calibri')
normal = document.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.03
normal.paragraph_format.space_after = Pt(5)
document.styles['Title'].font.size = Pt(23)
document.styles['Title'].paragraph_format.space_after = Pt(3)
document.styles['Subtitle'].font.size = Pt(12)
document.styles['Subtitle'].paragraph_format.space_after = Pt(6)
heading = document.styles['Heading 1']
heading.font.size = Pt(12)
heading.font.bold = True
heading.paragraph_format.space_before = Pt(10)
heading.paragraph_format.space_after = Pt(5)
heading.paragraph_format.keep_with_next = True
bullet = document.styles['List Bullet']
bullet.font.size = Pt(11)
bullet.paragraph_format.left_indent = Inches(0.16)
bullet.paragraph_format.first_line_indent = Inches(-0.1)
bullet.paragraph_format.space_after = Pt(3)

for block in source['blocks']:
    kind, text = block['kind'], block['text']
    style = {'title': 'Title', 'subtitle': 'Subtitle', 'heading': 'Heading 1',
             'bullet': 'List Bullet'}.get(kind, 'Normal')
    paragraph = document.add_paragraph(text, style=style)
    properties = paragraph._p.get_or_add_pPr()
    for border in list(properties.findall(qn('w:pBdr'))):
        properties.remove(border)
    if block.get('pageBreak'):
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.space_before = Pt(0)
    if kind == 'contact':
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
    if kind in ('employer', 'role'):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(7 if kind == 'role' else 0)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.bold = True

footer = section.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.add_run('Stephen M Abbott  |  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
document.core_properties.title = 'Stephen M Abbott Resume'
document.core_properties.author = 'Stephen M Abbott'
document.core_properties.subject = 'Cybersecurity and AI Executive'
document.core_properties.comments = 'Regenerated from verified approved master resume text.'
output_dir = ROOT / 'public/resume-files'
output_dir.mkdir(parents=True, exist_ok=True)
docx_output = output_dir / 'Stephen-Abbott-Resume.docx'
document.save(docx_output)
print(f'Created {docx_output.name}')

font_regular = 'Times-Roman'
font_bold = 'Times-Bold'
for regular, bold in (
    (
        Path('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'),
        Path('/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'),
    ),
    (
        Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
        Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    ),
):
    if regular.is_file() and bold.is_file():
        pdfmetrics.registerFont(TTFont('ResumeSans', str(regular)))
        pdfmetrics.registerFont(TTFont('ResumeSans-Bold', str(bold)))
        font_regular, font_bold = 'ResumeSans', 'ResumeSans-Bold'
        break

styles = {
    'title': ParagraphStyle(
        'ResumeTitle', fontName=font_bold, fontSize=23, leading=26,
        spaceAfter=4, alignment=TA_LEFT,
    ),
    'subtitle': ParagraphStyle(
        'ResumeSubtitle', fontName=font_regular, fontSize=12, leading=14,
        spaceAfter=6, alignment=TA_LEFT,
    ),
    'contact': ParagraphStyle(
        'ResumeContact', fontName=font_regular, fontSize=9.5, leading=12,
        spaceAfter=3, alignment=TA_LEFT,
    ),
    'heading': ParagraphStyle(
        'ResumeHeading', fontName=font_bold, fontSize=12, leading=14,
        spaceBefore=10, spaceAfter=5, alignment=TA_LEFT,
    ),
    'paragraph': ParagraphStyle(
        'ResumeParagraph', fontName=font_regular, fontSize=11, leading=13.3,
        spaceAfter=5, alignment=TA_LEFT,
    ),
    'employer': ParagraphStyle(
        'ResumeEmployer', fontName=font_bold, fontSize=11, leading=13,
        spaceAfter=3, alignment=TA_LEFT,
    ),
    'role': ParagraphStyle(
        'ResumeRole', fontName=font_bold, fontSize=11, leading=13,
        spaceBefore=7, spaceAfter=3, alignment=TA_LEFT,
    ),
    'bullet': ParagraphStyle(
        'ResumeBullet', fontName=font_regular, fontSize=11, leading=13.3,
        leftIndent=14, spaceAfter=3, alignment=TA_LEFT,
    ),
}

story = []
for block in source['blocks']:
    kind, text = block['kind'], escape(block['text'])
    if block.get('pageBreak'):
        story.append(PageBreak())
    style = styles.get(kind, styles['paragraph'])
    prefix = '• ' if kind == 'bullet' else ''
    story.append(Paragraph(f'{prefix}{text}', style))


def add_page_number(canvas, _doc):
    canvas.saveState()
    canvas.setFont(font_regular, 8)
    canvas.drawString(0.65 * inch, 0.35 * inch, f'Stephen M Abbott  |  {canvas.getPageNumber()}')
    canvas.restoreState()


pdf_output = output_dir / 'Stephen-Abbott-Resume.pdf'
pdf = SimpleDocTemplate(
    str(pdf_output),
    pagesize=letter,
    leftMargin=0.65 * inch,
    rightMargin=0.65 * inch,
    topMargin=0.65 * inch,
    bottomMargin=0.65 * inch,
    title='Stephen M Abbott Resume',
    author='Stephen M Abbott',
    subject='Cybersecurity and AI Executive',
)
pdf.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
print(f'Created {pdf_output.name}')
